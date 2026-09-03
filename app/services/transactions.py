import csv
import io
import re
from datetime import datetime
from typing import Any, NamedTuple

from openpyxl import load_workbook

from app.schemas import Transaction

DATE_HEADERS = ("거래일시", "거래일자", "거래일", "일시", "날짜", "일자", "date", "datetime")
TIME_HEADERS = ("거래시간", "시간", "time")
IN_HEADERS = ("입금액", "입금", "맡기신금액", "맡기신 금액", "credit", "deposit")
OUT_HEADERS = ("출금액", "출금", "찾으신금액", "찾으신 금액", "debit", "withdrawal")
AMOUNT_HEADERS = ("거래금액", "금액", "amount")
DIRECTION_HEADERS = ("구분", "입출금구분", "거래구분", "type", "direction")
PARTY_HEADERS = ("거래처", "상대방", "보낸분", "받는분", "보낸분/받는분", "받는분/보낸분", "상대계좌", "counterparty", "name")
MEMO_HEADERS = ("적요", "내용", "거래내용", "메모", "비고", "거래기록사항", "memo", "description")

SPEND_KEYWORDS = (
    "카드", "체크", "결제", "atm", "현금", "인출", "자동이체", "통신", "공과금", "전기",
    "가스", "수도", "관리비", "임대료", "월세", "보험", "이자", "대출", "적금", "예금",
    "급여", "세금", "국민연금", "건강보험", "배달", "편의점", "마트", "cu", "gs25",
    "쿠팡", "네이버페이", "카카오페이", "토스", "정기", "납부", "요금",
)

DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y.%m.%d %H:%M:%S",
    "%Y.%m.%d %H:%M", "%Y.%m.%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d",
    "%Y%m%d %H:%M:%S", "%Y%m%d%H%M%S", "%Y%m%d", "%y.%m.%d %H:%M", "%y.%m.%d",
    "%m/%d/%Y %H:%M", "%m/%d/%Y",
)


class ParseError(ValueError):
    pass


class ParseResult(NamedTuple):
    transactions: list[Transaction]
    skipped_rows: int


def classify_kind(text: str) -> str:
    lowered = text.lower()
    if any(k in lowered for k in SPEND_KEYWORDS):
        return "spend"
    return "transfer"


def _norm(s: Any) -> str:
    return re.sub(r"\s+", "", str(s or "")).lower()


def _match(header: str, candidates: tuple[str, ...]) -> bool:
    h = _norm(header)
    return any(_norm(c) == h or _norm(c) in h for c in candidates)


def _parse_amount(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    s = str(value).strip().replace(",", "").replace("원", "").replace(" ", "")
    if not s or s in {"-", "0"}:
        return 0 if s == "0" else None
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    s = s.strip("-()+")
    try:
        n = int(round(float(s)))
    except ValueError:
        return None
    return -n if neg else n


def _parse_datetime(value: Any, time_value: Any = None) -> datetime | None:
    if isinstance(value, datetime):
        return value
    s = str(value or "").strip()
    if not s:
        return None
    if time_value:
        t = str(time_value).strip()
        if isinstance(time_value, datetime):
            t = time_value.strftime("%H:%M:%S")
        s = f"{s} {t}"
    s = re.sub(r"\s+", " ", s)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    m = re.match(r"(\d{4})[.\-/]?(\d{1,2})[.\-/]?(\d{1,2})(?:\D+(\d{1,2}):(\d{2})(?::(\d{2}))?)?", s)
    if m:
        y, mo, d, hh, mm, ss = m.groups()
        try:
            return datetime(int(y), int(mo), int(d), int(hh or 0), int(mm or 0), int(ss or 0))
        except ValueError:
            return None
    return None


def _decode(body: bytes) -> str:
    for enc in ("utf-8-sig", "cp949", "euc-kr", "utf-16"):
        try:
            return body.decode(enc)
        except UnicodeDecodeError:
            continue
    return body.decode("utf-8", errors="replace")


def _rows_from_csv(body: bytes) -> list[list[Any]]:
    text = _decode(body)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _rows_from_xlsx(body: bytes) -> list[list[Any]]:
    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _find_header(rows: list[list[Any]]) -> int:
    for i, row in enumerate(rows[:30]):
        cells = [str(c or "") for c in row]
        hits = sum(
            1
            for c in cells
            if _match(c, DATE_HEADERS + IN_HEADERS + OUT_HEADERS + AMOUNT_HEADERS + PARTY_HEADERS + MEMO_HEADERS)
        )
        if hits >= 2:
            return i
    raise ParseError("거래내역 헤더(거래일시·입금액·출금액 등)를 찾지 못했습니다.")


def _column_map(header: list[Any]) -> dict[str, int]:
    cols: dict[str, int] = {}
    for idx, cell in enumerate(header):
        name = str(cell or "")
        if not name:
            continue
        if "date" not in cols and _match(name, DATE_HEADERS):
            cols["date"] = idx
        elif "time" not in cols and _match(name, TIME_HEADERS):
            cols["time"] = idx
        elif "in" not in cols and _match(name, IN_HEADERS):
            cols["in"] = idx
        elif "out" not in cols and _match(name, OUT_HEADERS):
            cols["out"] = idx
        elif "amount" not in cols and _match(name, AMOUNT_HEADERS):
            cols["amount"] = idx
        elif "direction" not in cols and _match(name, DIRECTION_HEADERS):
            cols["direction"] = idx
        elif "party" not in cols and _match(name, PARTY_HEADERS):
            cols["party"] = idx
        elif "memo" not in cols and _match(name, MEMO_HEADERS):
            cols["memo"] = idx
    if "date" not in cols:
        raise ParseError("거래일시 컬럼을 찾지 못했습니다.")
    if "in" not in cols and "out" not in cols and "amount" not in cols:
        raise ParseError("입금액·출금액 또는 금액 컬럼을 찾지 못했습니다.")
    return cols


def _cell(row: list[Any], idx: int | None) -> Any:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _row_to_transaction(row: list[Any], cols: dict[str, int]) -> Transaction | None:
    when = _parse_datetime(_cell(row, cols.get("date")), _cell(row, cols.get("time")))
    if when is None:
        return None
    party = str(_cell(row, cols.get("party")) or "").strip()
    memo = str(_cell(row, cols.get("memo")) or "").strip()
    amount: int | None = None
    direction: str | None = None
    if "in" in cols or "out" in cols:
        inc = _parse_amount(_cell(row, cols.get("in"))) or 0
        outc = _parse_amount(_cell(row, cols.get("out"))) or 0
        if inc > 0:
            amount, direction = inc, "in"
        elif outc > 0:
            amount, direction = outc, "out"
    if amount is None and "amount" in cols:
        raw = _parse_amount(_cell(row, cols.get("amount")))
        if raw is None or raw == 0:
            return None
        dir_text = _norm(_cell(row, cols.get("direction")))
        if dir_text:
            direction = "in" if any(k in dir_text for k in ("입금", "입", "in", "credit", "+")) else "out"
        else:
            direction = "in" if raw > 0 else "out"
        amount = abs(raw)
    if amount is None or not direction or amount <= 0:
        return None
    kind = classify_kind(f"{party} {memo}") if direction == "out" else "unknown"
    return Transaction(
        occurred_at=when,
        amount=amount,
        direction=direction,
        counterparty=party or memo,
        memo=memo,
        kind=kind,
    )


def rows_to_transactions(rows: list[list[Any]]) -> ParseResult:
    start = _find_header(rows)
    cols = _column_map(rows[start])
    out: list[Transaction] = []
    skipped = 0
    for row in rows[start + 1 :]:
        if not row or all(c in (None, "") for c in row):
            continue
        try:
            tx = _row_to_transaction(row, cols)
        except Exception:
            skipped += 1
            continue
        if tx is None:
            skipped += 1
            continue
        out.append(tx)
    if not out:
        raise ParseError("파싱된 거래가 없습니다. 날짜·금액 형식을 확인해 주세요.")
    out.sort(key=lambda t: t.occurred_at)
    return ParseResult(out, skipped)


def parse_transactions(filename: str, body: bytes) -> ParseResult:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix in ("xlsx", "xlsm"):
        rows = _rows_from_xlsx(body)
    elif suffix == "xls":
        raise ParseError("xls(구형 엑셀)는 지원하지 않습니다. xlsx 또는 CSV로 저장해 주세요.")
    else:
        rows = _rows_from_csv(body)
    return rows_to_transactions(rows)
